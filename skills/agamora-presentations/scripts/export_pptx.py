# -*- coding: utf-8 -*-
"""HTML-дек → редактируемый PPTX с той же геометрией, что у PDF.

Как работает:
1. В копию дека вставляется deck_export.js; headless Chrome (--dump-dom) раскладывает каждый слайд
   на родные элементы: прямоугольники и скругления, линии со стрелками, многоугольники и кривые SVG,
   картинки и текстовые блоки (с точными строками, шрифтом, кеглем, цветом).
2. То, что нельзя честно выразить фигурами PowerPoint (фильтры, маски, сложные тени, текст по кривой),
   становится «растровым островом»: Chrome печатает его отдельной страницей, остров вставляется картинкой
   на своё место и в свой порядок наложения.
3. python-pptx собирает слайды 13,333×7,5 дюйма (1 px = 0,6 pt): текст остаётся редактируемым, переносы
   строк — как в PDF, заметки докладчика из <aside class="notes">.

Запуск:  python export_pptx.py deck.html [deck.pptx] [--font "IBM Plex Sans=Segoe UI"] [--reflow] [--keep]
  --font    подменить шрифт в PPTX (если у получателя не установлен встроенный в HTML шрифт);
  --reflow  не ставить принудительные переносы — PowerPoint переносит строки сам;
  --keep    не удалять временные файлы (для разбора).
"""
import argparse
import base64
import io
import json
import math
import os
import re
import shutil
import sys
import tempfile
import urllib.parse

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import chrome  # noqa: E402

SKILL = os.path.dirname(HERE)
EMU_PX = 7620            # 12 192 000 EMU / 1600 px
PT_PX = 0.6              # 1 px = 0,6 pt
BASE_RATIO = 0.8         # PowerPoint при точном межстрочном ставит базовую линию на 0,8 высоты строки
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Начертания: CSS-насыщенность → имя семейства в PowerPoint (устаревшие имена GDI) и признак «жирный».
FACES = {
    "IBM Plex Sans": {400: ("IBM Plex Sans", False), 500: ("IBM Plex Sans Medm", False),
                      600: ("IBM Plex Sans SmBld", False), 700: ("IBM Plex Sans", True)},
    "Segoe UI": {300: ("Segoe UI Light", False), 350: ("Segoe UI Semilight", False), 400: ("Segoe UI", False),
                 600: ("Segoe UI Semibold", False), 700: ("Segoe UI", True), 900: ("Segoe UI Black", False)},
    "Onest": {400: ("Onest", False), 500: ("Onest Medium", False), 600: ("Onest SemiBold", False), 700: ("Onest", True)},
    "Unbounded": {500: ("Unbounded Medium", False), 600: ("Unbounded SemiBold", False), 700: ("Unbounded", True)},
    "Caveat": {400: ("Caveat", True), 700: ("Caveat", True)},
    "Segoe Print": {400: ("Segoe Print", False), 700: ("Segoe Print", True)},
}
# Файлы шрифтов для проверки «влезает ли строка» (метрики того шрифта, что будет в PPTX).
FONT_FILES = {
    ("IBM Plex Sans", 400): "assets/fonts/ibm-plex-sans/IBMPlexSans-Regular.ttf",
    ("IBM Plex Sans", 500): "assets/fonts/ibm-plex-sans/IBMPlexSans-Medium.ttf",
    ("IBM Plex Sans", 600): "assets/fonts/ibm-plex-sans/IBMPlexSans-SemiBold.ttf",
    ("IBM Plex Sans", 700): "assets/fonts/ibm-plex-sans/IBMPlexSans-Bold.ttf",
    ("Caveat", 700): "assets/fonts/caveat/Caveat-Bold.ttf",
    ("Onest", 400): "assets/fonts/onest/Onest-Regular.ttf",
    ("Onest", 500): "assets/fonts/onest/Onest-Medium.ttf",
    ("Onest", 600): "assets/fonts/onest/Onest-SemiBold.ttf",
    ("Onest", 700): "assets/fonts/onest/Onest-Bold.ttf",
    ("Unbounded", 500): "assets/fonts/unbounded/Unbounded-Medium.ttf",
    ("Unbounded", 600): "assets/fonts/unbounded/Unbounded-SemiBold.ttf",
    ("Unbounded", 700): "assets/fonts/unbounded/Unbounded-Bold.ttf",
    ("Segoe UI", 400): "C:/Windows/Fonts/segoeui.ttf",
    ("Segoe UI", 600): "C:/Windows/Fonts/seguisb.ttf",
    ("Segoe UI", 700): "C:/Windows/Fonts/segoeuib.ttf",
}


def E(v):
    return Emu(int(round(v * EMU_PX)))


def a(tag):
    return "{%s}%s" % (A_NS, tag)


def sub(parent, tag, **attrs):
    el = etree.SubElement(parent, a(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


# ── цвет и линии ───────────────────────────────────────────────────────────────────────────────
def solid(parent, col):
    """<a:solidFill> с прозрачностью; col = {"c": "RRGGBB", "a": 0..1}."""
    f = sub(parent, "solidFill")
    c = sub(f, "srgbClr", val=col["c"])
    if col.get("a", 1) < 0.999:
        sub(c, "alpha", val=int(round(col["a"] * 100000)))
    return f


def strip_style(sp):
    st = sp._element.find(qn("p:style"))
    if st is not None:
        sp._element.remove(st)


def sppr(sp):
    return sp._element.spPr


def set_fill(sp, fill, grad=None):
    pr = sppr(sp)
    for tag in ("a:noFill", "a:solidFill", "a:gradFill"):
        for el in pr.findall(qn(tag)):
            pr.remove(el)
    geom = pr.find(qn("a:prstGeom"))
    if geom is None:
        geom = pr.find(qn("a:custGeom"))
    idx = list(pr).index(geom) + 1 if geom is not None else len(pr)
    if grad:
        g = etree.Element(a("gradFill"), rotWithShape="1")
        gs = sub(g, "gsLst")
        for st in grad["stops"]:
            s = sub(gs, "gs", pos=int(round(st["pos"] * 100000)))
            c = sub(s, "srgbClr", val=st["c"])
            if st.get("a", 1) < 0.999:
                sub(c, "alpha", val=int(round(st["a"] * 100000)))
        sub(g, "lin", ang=int(round(((grad["ang"] - 90) % 360) * 60000)), scaled="0")
        pr.insert(idx, g)
    elif fill:
        f = etree.Element(a("solidFill"))
        c = sub(f, "srgbClr", val=fill["c"])
        if fill.get("a", 1) < 0.999:
            sub(c, "alpha", val=int(round(fill["a"] * 100000)))
        pr.insert(idx, f)
    else:
        pr.insert(idx, etree.Element(a("noFill")))


def set_line(sp, line, head=None, tail=None):
    pr = sppr(sp)
    for el in pr.findall(qn("a:ln")):
        pr.remove(el)
    ln = etree.Element(a("ln"))
    if not line:
        sub(ln, "noFill")
    else:
        ln.set("w", str(int(round(line["w"] * EMU_PX))))
        cap = {"round": "rnd", "square": "sq"}.get(line.get("cap"), None)
        if cap:
            ln.set("cap", cap)
        solid(ln, line)
        da = line.get("dashArr")
        if da and line["w"] > 0:
            cd = sub(ln, "custDash")
            vals = da if len(da) % 2 == 0 else da * 2
            for i in range(0, len(vals), 2):
                sub(cd, "ds", d=max(1, int(round(vals[i] / line["w"] * 100000))),
                    sp=max(1, int(round(vals[i + 1] / line["w"] * 100000))))
        elif line.get("dash"):
            sub(ln, "prstDash", val=line["dash"])
        if line.get("join") == "round":
            sub(ln, "round")
        for tag, mk in (("headEnd", head), ("tailEnd", tail)):
            if mk:
                sub(ln, tag, type=mk.get("type", "triangle"), w=mk.get("size", "med"), len=mk.get("size", "med"))
    # <a:ln> идёт после заливки и перед эффектами
    eff = pr.find(qn("a:effectLst"))
    if eff is not None:
        eff.addprevious(ln)
    else:
        pr.append(ln)


def set_shadow(sp, sh):
    if not sh:
        return
    pr = sppr(sp)
    eff = sub(pr, "effectLst")
    dist = math.hypot(sh["x"], sh["y"])
    ang = (math.degrees(math.atan2(sh["y"], sh["x"])) % 360) if dist else 0
    o = sub(eff, "outerShdw", blurRad=int(sh["blur"] * EMU_PX), dist=int(dist * EMU_PX),
            dir=int(ang * 60000), algn="ctr", rotWithShape="0")
    c = sub(o, "srgbClr", val=sh["c"])
    sub(c, "alpha", val=int(round(sh["a"] * 100000)))


# ── фигуры ─────────────────────────────────────────────────────────────────────────────────────
GEOM = {"rect": MSO_SHAPE.RECTANGLE, "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
        "round2SameRect": MSO_SHAPE.ROUND_2_SAME_RECTANGLE}


def add_box(slide, it):
    w, h = max(it["w"], 0.5), max(it["h"], 0.5)
    sp = slide.shapes.add_shape(GEOM.get(it["geom"], MSO_SHAPE.RECTANGLE), E(it["x"]), E(it["y"]), E(w), E(h))
    strip_style(sp)
    if it.get("adj"):
        for i, v in enumerate(it["adj"]):
            if i < len(sp.adjustments):
                sp.adjustments[i] = max(0.0, min(0.5, v))
    set_fill(sp, it.get("fill"), it.get("grad"))
    set_line(sp, it.get("line"))
    set_shadow(sp, it.get("shadow"))
    if it.get("rot"):
        sp.rotation = it["rot"]
    return sp


def add_line(slide, it):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(it["x1"]), E(it["y1"]), E(it["x2"]), E(it["y2"]))
    strip_style(c)
    set_line(c, it["line"], it.get("head"), it.get("tail"))
    return c


def add_custom(slide, segs, fill, line, head=None, tail=None):
    pts = [p for s in segs for p in s["p"]]
    if not pts:
        return None
    x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
    x1, y1 = max(p[0] for p in pts), max(p[1] for p in pts)
    w, h = max(x1 - x0, 0.5), max(y1 - y0, 0.5)
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x0), E(y0), E(w), E(h))
    strip_style(sp)
    pr = sppr(sp)
    old = pr.find(qn("a:prstGeom"))
    cg = etree.Element(a("custGeom"))
    for t in ("avLst", "gdLst", "ahLst", "cxnLst"):
        sub(cg, t)
    sub(cg, "rect", l="0", t="0", r="r", b="b")
    pl = sub(cg, "pathLst")
    W, H = int(round(w * EMU_PX)), int(round(h * EMU_PX))
    path = sub(pl, "path", w=W, h=H)
    if not fill:
        path.set("fill", "none")

    def P(el, p):
        sub(el, "pt", x=int(round((p[0] - x0) * EMU_PX)), y=int(round((p[1] - y0) * EMU_PX)))

    for s in segs:
        i = 0
        for cmd in s["c"]:
            if cmd == "M":
                P(sub(path, "moveTo"), s["p"][i]); i += 1
            elif cmd == "L":
                P(sub(path, "lnTo"), s["p"][i]); i += 1
            elif cmd == "C":
                cb = sub(path, "cubicBezTo")
                for k in range(3):
                    P(cb, s["p"][i + k])
                i += 3
        if s.get("closed"):
            sub(path, "close")
    old.addprevious(cg)
    pr.remove(old)
    set_fill(sp, fill)
    set_line(sp, line, head, tail)
    return sp


def add_poly(slide, it):
    seg = {"p": it["pts"], "c": ["M"] + ["L"] * (len(it["pts"]) - 1), "closed": it.get("closed")}
    return add_custom(slide, [seg], it.get("fill"), it.get("line"), it.get("head"), it.get("tail"))


def add_path(slide, it):
    return add_custom(slide, it["segs"], it.get("fill"), it.get("line"), it.get("head"), it.get("tail"))


def image_bytes(src, base_dir):
    if src.startswith("data:"):
        head, data = src.split(",", 1)
        return base64.b64decode(data) if ";base64" in head else urllib.parse.unquote(data).encode()
    if src.startswith("file:"):
        path = urllib.parse.unquote(urllib.parse.urlparse(src).path)
        if re.match(r"^/[A-Za-z]:", path):
            path = path[1:]
        return open(path, "rb").read()
    p = os.path.join(base_dir, src)
    return open(p, "rb").read() if os.path.exists(p) else None


def add_img(slide, it, base_dir, warn):
    data = image_bytes(it["src"], base_dir)
    if not data:
        warn.append("картинка не найдена: " + it["src"][:80])
        return None
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    nw, nh = it.get("nw") or w, it.get("nh") or h
    crop = None
    if it.get("fit") == "contain" and nw and nh:
        k = min(w / nw, h / nh)
        cw, ch = nw * k, nh * k
        x, y, w, h = x + (w - cw) / 2, y + (h - ch) / 2, cw, ch
    elif it.get("fit") == "cover" and nw and nh:
        k = max(w / nw, h / nh)
        cw, ch = nw * k, nh * k
        crop = ((cw - w) / 2 / cw, (ch - h) / 2 / ch)
    pic = slide.shapes.add_picture(io.BytesIO(data), E(x), E(y), E(w), E(h))
    if crop:
        pic.crop_left = pic.crop_right = crop[0]
        pic.crop_top = pic.crop_bottom = crop[1]
    if it.get("radius"):
        g = pic._element.spPr.find(qn("a:prstGeom"))
        g.set("prst", "roundRect")
        av = g.find(qn("a:avLst"))
        sub(av, "gd", name="adj", fmla="val %d" % int(it["radius"] * 100000))
    if it.get("alpha"):
        blip = pic._element.find(".//" + qn("a:blip"))
        sub(blip, "alphaModFix", amt=int(it["alpha"] * 100000))
    return pic


# ── текст ──────────────────────────────────────────────────────────────────────────────────────
class Fonts:
    def __init__(self, font_map):
        self.map = font_map
        self.cache = {}

    def family(self, fam):
        return self.map.get(fam, fam)

    def face(self, fam, weight):
        fam = self.family(fam)
        t = FACES.get(fam)
        if not t:
            return fam, weight >= 600
        k = min(t, key=lambda w: (abs(w - weight), -w))
        return t[k]

    def pil(self, fam, weight, size):
        from PIL import ImageFont
        fam = self.family(fam)
        cands = [k for k in FONT_FILES if k[0] == fam]
        if not cands:
            return None
        k = min(cands, key=lambda c: abs(c[1] - weight))
        path = FONT_FILES[k] if os.path.isabs(FONT_FILES[k]) or FONT_FILES[k].startswith("C:") else os.path.join(SKILL, FONT_FILES[k])
        if not os.path.exists(path):
            return None
        key = (path, round(size * 4))
        if key not in self.cache:
            self.cache[key] = ImageFont.truetype(path, max(1, int(round(size))))
        return self.cache[key]


def split_lines(it):
    """Разрезает серии текста по строкам, как их разложил Chrome. Возвращает [[(текст, стиль), …], …]
    или None, если сопоставить не удалось (тогда PowerPoint переносит строки сам)."""
    runs = it["runs"]
    flat, styles = "", []
    for r in runs:
        if r.get("br"):
            flat += "\n"; styles.append(None)
        else:
            flat += r["t"]; styles.extend([r["s"]] * len(r["t"]))
    breaks, pos = [], 0
    for ln in it["lines"][:-1]:
        t = ln["text"].strip()
        while pos < len(flat) and flat[pos] in " \n":
            pos += 1
        if not flat.startswith(t, pos):
            return None
        pos += len(t)
        if ln.get("br") and pos < len(flat) and flat[pos] == "\n":
            breaks.append((pos, pos + 1)); pos += 1
        else:
            end = pos
            while end < len(flat) and flat[end] == " ":
                end += 1
            breaks.append((pos, end)); pos = end
    out, start = [], 0
    for b0, b1 in breaks + [(len(flat), len(flat))]:
        seg, cur, cur_st = [], "", None
        for ch, st in zip(flat[start:b0], styles[start:b0]):
            if ch == "\n":
                continue
            if st is not cur_st and cur:
                seg.append((cur, cur_st)); cur = ""
            cur_st = st; cur += ch
        if cur:
            seg.append((cur, cur_st))
        out.append(seg)
        start = b1
    return out


def rpr(parent, st, fonts, tag="rPr", scale=1.0):
    r = sub(parent, tag, lang="ru-RU", dirty="0")
    r.set("sz", str(int(round(st["size"] * PT_PX * 100 * scale))))
    name, bold = fonts.face(st["font"], st.get("weight", 400))
    if bold:
        r.set("b", "1")
    if st.get("italic"):
        r.set("i", "1")
    if st.get("u"):
        r.set("u", "sng")
    if st.get("s"):
        r.set("strike", "sngStrike")
    if st.get("caps") == "all":
        r.set("cap", "all")
    elif st.get("caps") == "small":
        r.set("cap", "small")
    if st.get("ls"):
        r.set("spc", str(int(round(st["ls"] * PT_PX * 100))))
    if st.get("va") == "super":
        r.set("baseline", "30000")
    elif st.get("va") == "sub":
        r.set("baseline", "-25000")
    if st.get("color"):
        solid(r, st["color"])
    sub(r, "latin", typeface=name)
    sub(r, "cs", typeface=name)
    return r


def add_text(slide, it, fonts, reflow, warn):
    runs = [r for r in it["runs"] if r.get("br") or r["t"]]
    if not runs:
        return None
    st0 = next(r["s"] for r in runs if not r.get("br"))
    lines = it["lines"]
    lh = it.get("lh") or st0["size"] * 1.25
    per_line = None if reflow else split_lines(it)
    # вертикаль: базовая линия первой строки как в Chrome
    base = lines[0]["base"]
    y = base - BASE_RATIO * lh
    h = max(lh * len(lines), lines[-1]["bottom"] - y) + 2
    # горизонталь: рамка содержимого блока с запасом на разницу метрик
    align = it.get("align", "l")
    x, w = it["x"], it["w"]
    if it.get("svg"):
        pass
    elif not it.get("wrap", True) or per_line is not None:
        used = max(l["right"] for l in lines) - min(l["left"] for l in lines)
        w = max(w, used)
    slack = max(6.0, w * 0.04)
    if align == "ctr":
        x, w = x - slack / 2, w + slack
    elif align == "r":
        x, w = x - slack, w + slack
    else:
        w = w + slack
    # проверка «влезает ли строка» метриками шрифта PPTX; при нехватке — ужимаем кегль до 8 %
    scale = 1.0
    if per_line:
        worst = 0.0
        for seg in per_line:
            tw = 0.0
            for t, st in seg:
                f = fonts.pil(st["font"], st.get("weight", 400), st["size"])
                if f is None:
                    tw = 0; break
                tw += f.getlength(t) + (st.get("ls") or 0) * len(t)
            worst = max(worst, tw / max(w, 1))
        if worst > 1.0:
            scale = max(0.92, 1 / worst)
            if worst > 1 / 0.92:
                warn.append("строка не помещается в PPTX даже с кеглем −8 %%: «%s»" % lines[0]["text"][:50])
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    body = tb.text_frame._txBody
    bp = body.find(qn("a:bodyPr"))
    for k, v in (("wrap", "square" if (it.get("wrap", True) and not it.get("svg")) else "none"),
                 ("lIns", "0"), ("tIns", "0"), ("rIns", "0"), ("bIns", "0"), ("anchor", "t"), ("rtlCol", "0")):
        bp.set(k, v)
    for ch in list(bp):
        bp.remove(ch)
    sub(bp, "noAutofit")
    for p_ in body.findall(qn("a:p")):
        body.remove(p_)
    p = sub(body, "p")
    ppr = sub(p, "pPr", algn=align)
    bullet = it.get("bullet")
    if bullet:
        ppr.set("marL", str(int(bullet["w"] * EMU_PX)))
        ppr.set("indent", str(-int(bullet["w"] * EMU_PX)))
    lnsp = sub(ppr, "lnSpc")
    sub(lnsp, "spcPts", val=int(round(lh * PT_PX * 100)))
    sub(sub(ppr, "spcBef"), "spcPts", val="0")
    sub(sub(ppr, "spcAft"), "spcPts", val="0")
    if bullet:
        if bullet.get("color"):
            sub(sub(ppr, "buClr"), "srgbClr", val=bullet["color"]["c"])
        if bullet.get("ch"):
            sub(ppr, "buChar", char=bullet["ch"])
        else:
            sub(ppr, "buAutoNum", type=bullet.get("auto", "arabicPeriod"), startAt=bullet.get("start", 1))
    else:
        sub(ppr, "buNone")

    def run(t, st):
        r = sub(p, "r")
        rpr(r, st, fonts, scale=scale)
        tt = sub(r, "t")
        tt.text = t.replace("\u2060", "")

    if per_line:
        for i, seg in enumerate(per_line):
            if i:
                rpr(sub(p, "br"), (seg[0][1] if seg else st0), fonts, scale=scale)
            for t, st in seg:
                run(t, st)
    else:
        for r in runs:
            if r.get("br"):
                rpr(sub(p, "br"), st0, fonts, scale=scale)
            else:
                run(r["t"].strip("\n"), r["s"])
    rpr(p, st0, fonts, "endParaRPr", scale=scale)
    if it.get("rot"):
        tb.rotation = it["rot"]
    return tb


# ── растровые острова ──────────────────────────────────────────────────────────────────────────
def island_images(raster_pdf, layout, tmp, zoom=2.0):
    """Кадры островов: страница PDF → PNG с прозрачностью → обрезка по непрозрачной области."""
    import pymupdf
    from PIL import Image
    doc = pymupdf.open(raster_pdf)
    res, k = {}, 0
    for si, s in enumerate(layout["slides"]):
        for ii, isl in enumerate(s["islands"]):
            if k >= doc.page_count:
                return res
            z = zoom * 1600 / doc[k].rect.width     # страница в пунктах: 1600 px = 1200 pt
            pix = doc[k].get_pixmap(matrix=pymupdf.Matrix(z, z), alpha=True)
            im = Image.frombytes("RGBA", (pix.width, pix.height), pix.samples)
            bbox = im.getchannel("A").getbbox()
            opaque = bbox == (0, 0, im.width, im.height) and isl.get("bb", {}).get("w", 1600) < 1590
            if not bbox:
                k += 1; continue
            if opaque:    # фон не прозрачный — режем по рамке элемента
                bb = isl["bb"]
                bbox = (int(bb["x"] * zoom), int(bb["y"] * zoom), int((bb["x"] + bb["w"]) * zoom + 1), int((bb["y"] + bb["h"]) * zoom + 1))
            crop = im.crop(bbox)
            path = os.path.join(tmp, "island_%02d_%02d.png" % (si, ii))
            crop.save(path)
            res[(si, ii)] = (path, bbox[0] / zoom, bbox[1] / zoom, crop.width / zoom, crop.height / zoom)
            k += 1
    return res


# ── сборка ─────────────────────────────────────────────────────────────────────────────────────
def inject(html_path):
    src = io.open(html_path, encoding="utf-8").read()
    js = io.open(os.path.join(HERE, "deck_export.js"), encoding="utf-8").read()
    tag = "<script>" + js + "</script>"
    out = src.replace("</body>", tag + "</body>") if "</body>" in src else src + tag
    d, b = os.path.split(os.path.abspath(html_path))
    tmp_html = os.path.join(d, "." + os.path.splitext(b)[0] + ".export.html")   # рядом с деком: относительные пути живы
    io.open(tmp_html, "w", encoding="utf-8").write(out)
    return tmp_html


def dump_layout(html_path, budget_ms=8000):
    """Раскладка дека (JSON) — общая для экспорта в PPTX и для проверки качества (qa.py)."""
    tmp_html = inject(html_path)
    m = None
    for attempt in range(3):                      # Chrome изредка отдаёт DOM до отработки сценария — повторяем
        dom = chrome.dump_dom(tmp_html, "export=layout", budget_ms * (attempt + 1))
        m = re.search(r'<script type="application/json" id="x-layout">({.*?)</script>', dom, re.S)
        if m:
            break
    if not m:
        raise RuntimeError("Chrome не вернул раскладку: проверьте, что дек открывается и в нём есть section.slide")
    return json.loads(m.group(1)), tmp_html


def export(html_path, out_path, font_map=None, reflow=False, keep=False, log=print):
    html_path = os.path.abspath(html_path)
    fonts = Fonts(font_map or {})
    tmp = tempfile.mkdtemp(prefix="deck-pptx-")
    layout, tmp_html = dump_layout(html_path)
    warn = list(layout.get("warnings") or [])
    try:
        islands = {}
        if any(s["islands"] for s in layout["slides"]):
            raster = chrome.print_pdf(tmp_html, os.path.join(tmp, "raster.pdf"), "export=raster")
            islands = island_images(raster, layout, tmp)
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        blank = prs.slide_layouts[6]
        base_dir = os.path.dirname(html_path)
        counts = {"text": 0, "shape": 0, "island": 0}
        for si, s in enumerate(layout["slides"]):
            slide = prs.slides.add_slide(blank)
            if s.get("bg"):
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor.from_string(s["bg"]["c"])
            for it in s["items"]:
                k = it["k"]
                try:
                    if k == "shape":
                        add_box(slide, it); counts["shape"] += 1
                    elif k == "line":
                        add_line(slide, it); counts["shape"] += 1
                    elif k == "poly":
                        add_poly(slide, it); counts["shape"] += 1
                    elif k == "path":
                        add_path(slide, it); counts["shape"] += 1
                    elif k == "text":
                        add_text(slide, it, fonts, reflow, warn); counts["text"] += 1
                    elif k == "img":
                        add_img(slide, it, base_dir, warn)
                    elif k == "island":
                        rec = islands.get((si, it["n"]))
                        if rec:
                            path, x, y, w, h = rec
                            slide.shapes.add_picture(path, E(x), E(y), E(w), E(h)); counts["island"] += 1
                        else:
                            warn.append("слайд %d: растровый остров %d не отрисован (%s)" % (si + 1, it["n"], it.get("why")))
                except Exception as e:  # один неудачный элемент не должен ронять весь дек
                    warn.append("слайд %d: элемент %s пропущен — %s" % (si + 1, k, e))
            if s.get("notes"):
                slide.notes_slide.notes_text_frame.text = s["notes"]
        prs.save(out_path)
        log("PPTX: %s — слайдов %d, текстовых блоков %d, фигур %d, растровых островов %d"
            % (out_path, len(layout["slides"]), counts["text"], counts["shape"], counts["island"]))
        for w_ in warn[:40]:
            log("  внимание: " + w_)
        return out_path, warn
    finally:
        if not keep:
            shutil.rmtree(tmp, ignore_errors=True)
            if os.path.exists(tmp_html):
                os.remove(tmp_html)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="HTML-дек → редактируемый PPTX")
    ap.add_argument("html")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--font", action="append", default=[], help='подмена шрифта: "IBM Plex Sans=Segoe UI"')
    ap.add_argument("--reflow", action="store_true", help="без принудительных переносов строк")
    ap.add_argument("--keep", action="store_true", help="оставить временные файлы")
    a_ = ap.parse_args()
    out = a_.out or os.path.splitext(a_.html)[0] + ".pptx"
    fm = dict(x.split("=", 1) for x in a_.font)
    export(a_.html, out, fm, a_.reflow, a_.keep)


if __name__ == "__main__":
    main()
